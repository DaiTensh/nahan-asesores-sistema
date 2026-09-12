# -*- coding: utf-8 -*-
"""RF35 — Exportando Reportes en Formato Excel.

GET /api/reportes/<id_reporte>/excel vuelve a consultar con los parámetros
guardados en `reporte` (misma capa que usaría el PDF de RF34, ver
_datos_para_exportar en reportes_routes.py) y arma un .xlsx: una fila por
tarea (o por usuario en productividad) y los encabezados del reporte.
"""
import io
import json

import pytest

from openpyxl import load_workbook

from tests.conftest import normalize_sql

import backend.routes.reportes_routes as reportes_routes


def _handler_tareas_por_cliente():
    reporte_guardado = {
        "id_reporte": 1,
        "tipo_reporte": "TAREAS_POR_CLIENTE",
        "parametros": json.dumps({
            "id_cliente": 10, "fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-31",
        }),
    }
    tareas = [
        {
            "id_tarea": 1, "titulo": "Tarea A", "estado": "PENDIENTE", "prioridad": "ALTA",
            "fecha_creacion": "01-01-2026 10:00", "fecha_vencimiento": "10-01-2026",
            "id_cliente": 10, "cliente": "Cliente Diez",
            "id_responsable": 2, "responsable": "Fulano",
        },
    ]

    def handler(sql, params, cursor):
        sql_normalizado = normalize_sql(sql)
        if "FROM reporte WHERE id_reporte" in sql_normalizado:
            return [reporte_guardado]
        if "FROM tarea t" in sql_normalizado and "INNER JOIN cliente c" in sql_normalizado:
            return [f.copy() for f in tareas]
        if "SELECT razon_social FROM cliente" in sql_normalizado:
            return [{"razon_social": "Cliente Diez"}]
        return []

    return handler


def _handler_productividad():
    reporte_guardado = {
        "id_reporte": 2,
        "tipo_reporte": "PRODUCTIVIDAD_POR_USUARIO",
        "parametros": json.dumps({
            "fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-31", "id_area": None,
        }),
    }
    usuarios = [{"id_usuario": 2, "nombres": "Fulano", "id_area": 1}]

    def handler(sql, params, cursor):
        sql_normalizado = normalize_sql(sql)
        if "FROM reporte WHERE id_reporte" in sql_normalizado:
            return [reporte_guardado]
        if "SELECT id_usuario, nombres, id_area FROM usuario" in sql_normalizado:
            return [u.copy() for u in usuarios]
        return []

    return handler


def _leer_xlsx(response):
    libro = load_workbook(io.BytesIO(response.data))
    return list(libro.active.iter_rows(values_only=True))


def test_solo_administrador_puede_exportar(client, iniciar_sesion):
    iniciar_sesion(usuario_id=1, rol_id=2)

    respuesta = client.get("/api/reportes/1/excel")

    assert respuesta.status_code == 403


def test_reporte_inexistente_responde_404(client, monkeypatch, iniciar_sesion, fake_connection_factory):
    iniciar_sesion(usuario_id=1, rol_id=1)

    connection = fake_connection_factory(lambda sql, params, cursor: [])
    monkeypatch.setattr(reportes_routes, "get_connection", lambda: connection)

    respuesta = client.get("/api/reportes/999/excel")

    assert respuesta.status_code == 404


def test_exporta_tareas_por_cliente_una_fila_por_tarea(client, monkeypatch, iniciar_sesion, fake_connection_factory):
    iniciar_sesion(usuario_id=1, rol_id=1)

    connection = fake_connection_factory(_handler_tareas_por_cliente())
    monkeypatch.setattr(reportes_routes, "get_connection", lambda: connection)

    respuesta = client.get("/api/reportes/1/excel")

    assert respuesta.status_code == 200
    assert respuesta.headers["Content-Type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "Cliente_Diez" in respuesta.headers["Content-Disposition"]
    assert "2026-01-01_a_2026-01-31" in respuesta.headers["Content-Disposition"]

    filas = _leer_xlsx(respuesta)
    assert filas[0] == ("Título", "Cliente", "Responsable", "Estado", "Prioridad",
                         "Fecha de creación", "Fecha de vencimiento")
    assert filas[1] == ("Tarea A", "Cliente Diez", "Fulano", "PENDIENTE", "ALTA",
                         "01-01-2026 10:00", "10-01-2026")
    assert len(filas) == 2


def test_exporta_productividad_una_fila_por_usuario(client, monkeypatch, iniciar_sesion, fake_connection_factory):
    iniciar_sesion(usuario_id=1, rol_id=1)

    connection = fake_connection_factory(_handler_productividad())
    monkeypatch.setattr(reportes_routes, "get_connection", lambda: connection)

    respuesta = client.get("/api/reportes/2/excel")

    assert respuesta.status_code == 200
    filas = _leer_xlsx(respuesta)
    assert filas[0] == ("Usuario", "Tareas completadas",
                         "Tiempo promedio de resolución (horas)", "Carga vigente")
    # openpyxl relee una celda escrita como "" como None: no es un valor real.
    assert filas[1] == ("Fulano", 0, None, 0)


def test_excel_no_ejecuta_formulas_del_usuario(client, monkeypatch, iniciar_sesion, fake_connection_factory):
    iniciar_sesion()
    original = _handler_tareas_por_cliente()
    formula = '=HYPERLINK("https://example.invalid")'

    def handler(sql, params, cursor):
        filas = original(sql, params, cursor)
        if filas and "titulo" in filas[0]:
            filas[0]["titulo"] = formula
        return filas

    conexion = fake_connection_factory(handler)
    monkeypatch.setattr(reportes_routes, "get_connection", lambda: conexion)
    respuesta = client.get("/api/reportes/1/excel")
    libro = load_workbook(io.BytesIO(respuesta.data))
    assert libro.active['A2'].value == formula
    assert libro.active['A2'].data_type == 's'
    assert conexion.closed


def test_pdf_abrible_con_texto_especial_y_varias_paginas(client, monkeypatch, iniciar_sesion, fake_connection_factory):
    from pypdf import PdfReader
    iniciar_sesion()
    original = _handler_tareas_por_cliente()

    def handler(sql, params, cursor):
        filas = original(sql, params, cursor)
        if filas and "titulo" in filas[0]:
            filas[0]["titulo"] = "Revisión Ñandú <legal> & contable"
            return [dict(filas[0], id_tarea=i) for i in range(100)]
        return filas

    conexion = fake_connection_factory(handler)
    monkeypatch.setattr(reportes_routes, "get_connection", lambda: conexion)
    respuesta = client.get("/api/reportes/1/pdf")
    assert respuesta.status_code == 200
    assert respuesta.mimetype == 'application/pdf'
    assert '.pdf' in respuesta.headers['Content-Disposition']
    pdf = PdfReader(io.BytesIO(respuesta.data), strict=True)
    assert len(pdf.pages) > 1
    texto = ' '.join(' '.join(p.extract_text() for p in pdf.pages).split())
    assert 'Revisión Ñandú <legal> & contable' in texto
    assert '2026-01-01 a 2026-01-31' in texto
    assert 'Fecha de generación' in texto
    assert 'PENDIENTE: 100' in texto
    assert conexion.closed


def test_pdf_productividad_y_reporte_vacio(client, monkeypatch, iniciar_sesion, fake_connection_factory):
    from pypdf import PdfReader
    iniciar_sesion()
    for handler, identificador in [(_handler_productividad(), 2), (_handler_tareas_por_cliente(), 1)]:
        def sin_tareas(sql, params, cursor):
            filas = handler(sql, params, cursor)
            return [] if filas and 'titulo' in filas[0] else filas
        conexion = fake_connection_factory(sin_tareas)
        monkeypatch.setattr(reportes_routes, "get_connection", lambda: conexion)
        respuesta = client.get(f"/api/reportes/{identificador}/pdf")
        assert respuesta.status_code == 200
        pdf = PdfReader(io.BytesIO(respuesta.data), strict=True)
        assert len(pdf.pages) == 1
        assert ('Fulano' if identificador == 2 else reportes_routes.MENSAJE_SIN_TAREAS) in pdf.pages[0].extract_text()


def test_pdf_requiere_admin_y_reporte_existente(client, monkeypatch, iniciar_sesion, fake_connection_factory):
    assert client.get('/api/reportes/1/pdf').status_code == 401
    iniciar_sesion(rol_id=2)
    assert client.get('/api/reportes/1/pdf').status_code == 403
    iniciar_sesion()
    conexion = fake_connection_factory(lambda *a: [])
    monkeypatch.setattr(reportes_routes, 'get_connection', lambda: conexion)
    assert client.get('/api/reportes/999/pdf').status_code == 404
    assert conexion.closed


@pytest.mark.parametrize('tipo,campo', [('cliente', 'id_cliente'), ('responsable', 'id_responsable')])
@pytest.mark.parametrize('con_tareas', [False, True])
def test_rf31_rf32_registran_solo_generacion_con_datos(client, monkeypatch, iniciar_sesion, fake_connection_factory, tipo, campo, con_tareas):
    iniciar_sesion()

    def handler(sql, params, cursor):
        if 'FROM cliente WHERE' in sql:
            return [{'id_cliente': 1, 'razon_social': 'Cliente', 'rut': '1-9'}]
        if 'FROM usuario WHERE' in sql:
            return [{'id_usuario': 1, 'nombres': 'Responsable', 'email': 'prueba@test.local'}]
        if 'FROM tarea t' in sql:
            assert params == ('2026-01-01', '2026-01-31', '1')
            return [{'id_tarea': 2, 'estado': 'PENDIENTE'}] if con_tareas else []
        if 'LAST_INSERT_ID()' in sql:
            return [{'id_reporte': 8, 'fecha_generacion': '31-01-2026 10:00'}]
        return []

    conexion = fake_connection_factory(handler)
    monkeypatch.setattr(reportes_routes, 'get_connection', lambda: conexion)
    respuesta = client.get(f'/api/reportes/tareas-por-{tipo}?{campo}=1&fecha_inicio=2026-01-01&fecha_fin=2026-01-31')
    assert respuesta.status_code == 200
    datos = respuesta.get_json()
    assert datos['resumen_por_estado'] == ({'PENDIENTE': 1} if con_tareas else {})
    assert bool(datos.get('id_reporte')) == con_tareas
    assert conexion.commits == int(con_tareas)
    assert any('INSERT INTO reporte' in sql for sql, _ in conexion.executed) == con_tareas
    assert conexion.closed


def test_exportacion_expone_nombre_archivo_al_frontend(client, monkeypatch, iniciar_sesion, fake_connection_factory):
    iniciar_sesion()
    conexion = fake_connection_factory(_handler_tareas_por_cliente())
    monkeypatch.setattr(reportes_routes, 'get_connection', lambda: conexion)
    respuesta = client.get('/api/reportes/1/excel', headers={'Origin': 'http://127.0.0.1:5500'})
    assert respuesta.status_code == 200
    assert respuesta.headers['Access-Control-Expose-Headers'] == 'Content-Disposition'

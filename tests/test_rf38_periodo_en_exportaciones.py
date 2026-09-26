# -*- coding: utf-8 -*-
"""RF38 / RF34 / RF35 — El período aplicado queda en el reporte y en los
archivos exportados.

- El atajo de período («Última semana», «Último mes», «Año actual») llega
  como `periodo_etiqueta`, se valida contra una lista blanca, se guarda en
  `reporte.parametros` junto con las fechas y vuelve en la respuesta.
- El PDF muestra período (con el atajo), filtros aplicados, fecha de
  generación del reporte y fecha de exportación.
- El Excel conserva la hoja «Reporte» (una fila por elemento, criterio de
  RF35) y agrega la hoja «Parámetros» con período, filtros y fechas: el
  archivo dice por sí mismo qué consulta representa.
"""
import io
import json

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

import backend.routes.reportes_routes as reportes_routes
from tests.conftest import normalize_sql


def _handler_generacion(capturado):
    def handler(sql, params, cursor):
        sql_n = normalize_sql(sql)
        if "FROM cliente WHERE id_cliente" in sql_n:
            return [{"id_cliente": 10, "rut": "1-9", "razon_social": "Cliente Diez"}]
        if "FROM tarea t" in sql_n:
            return [{"id_tarea": 1, "estado": "PENDIENTE"}]
        if "INSERT INTO reporte" in sql_n:
            capturado["parametros"] = json.loads(params[2])
        if "LAST_INSERT_ID()" in sql_n:
            return [{"id_reporte": 5, "fecha_generacion": "25-09-2026 10:00"}]
        return []
    return handler


def _handler_exportacion(parametros, tipo="TAREAS_POR_CLIENTE"):
    from datetime import datetime

    def handler(sql, params, cursor):
        sql_n = normalize_sql(sql)
        if "FROM reporte WHERE id_reporte" in sql_n:
            return [{
                "id_reporte": 5, "tipo_reporte": tipo,
                "parametros": json.dumps(parametros),
                "fecha_generacion": datetime(2026, 9, 1, 8, 30),
            }]
        if "FROM tarea t" in sql_n:
            return [{
                "id_tarea": 1, "titulo": "Tarea A", "estado": "PENDIENTE", "prioridad": "ALTA",
                "fecha_creacion": "01-09-2026 10:00", "fecha_vencimiento": "10-09-2026",
                "id_cliente": 10, "cliente": "Cliente Diez", "id_responsable": 2, "responsable": "Fulano",
            }]
        if "SELECT razon_social FROM cliente" in sql_n:
            return [{"razon_social": "Cliente Diez"}]
        if "SELECT nombre_area FROM area" in sql_n:
            return [{"nombre_area": "CONTABLE"}]
        if "SELECT id_usuario, nombres, id_area FROM usuario" in sql_n:
            return [{"id_usuario": 2, "nombres": "Fulano", "id_area": 2}]
        return []
    return handler


def test_el_atajo_de_periodo_se_guarda_en_parametros_y_vuelve_en_la_respuesta(
    client, monkeypatch, iniciar_sesion, fake_connection_factory
):
    iniciar_sesion()
    capturado = {}
    conexion = fake_connection_factory(_handler_generacion(capturado))
    monkeypatch.setattr(reportes_routes, "get_connection", lambda: conexion)

    datos = client.get(
        "/api/reportes/tareas-por-cliente?id_cliente=10"
        "&fecha_inicio=2026-08-25&fecha_fin=2026-09-25&periodo_etiqueta=mes"
    ).get_json()

    assert datos["periodo"] == {
        "fecha_inicio": "2026-08-25", "fecha_fin": "2026-09-25", "etiqueta": "Último mes",
    }
    assert capturado["parametros"] == {
        "id_cliente": 10, "fecha_inicio": "2026-08-25", "fecha_fin": "2026-09-25",
        "periodo_etiqueta": "mes",
    }


def test_sin_atajo_no_se_agrega_etiqueta(client, monkeypatch, iniciar_sesion, fake_connection_factory):
    iniciar_sesion()
    capturado = {}
    conexion = fake_connection_factory(_handler_generacion(capturado))
    monkeypatch.setattr(reportes_routes, "get_connection", lambda: conexion)

    datos = client.get(
        "/api/reportes/tareas-por-cliente?id_cliente=10&fecha_inicio=2026-08-25&fecha_fin=2026-09-25"
    ).get_json()

    assert datos["periodo"] == {"fecha_inicio": "2026-08-25", "fecha_fin": "2026-09-25"}
    assert "periodo_etiqueta" not in capturado["parametros"]


@pytest.mark.parametrize("endpoint", [
    "tareas-por-cliente?id_cliente=10",
    "tareas-por-responsable?id_responsable=2",
    "productividad?",
])
def test_un_atajo_desconocido_se_rechaza_antes_de_consultar(client, iniciar_sesion, endpoint):
    iniciar_sesion()

    respuesta = client.get(
        f"/api/reportes/{endpoint}&fecha_inicio=2026-08-25&fecha_fin=2026-09-25"
        "&periodo_etiqueta=<script>"
    )

    assert respuesta.status_code == 400


def test_el_rango_invertido_se_rechaza(client, iniciar_sesion):
    iniciar_sesion()

    respuesta = client.get(
        "/api/reportes/tareas-por-cliente?id_cliente=10&fecha_inicio=2026-09-25&fecha_fin=2026-08-25"
    )

    assert respuesta.status_code == 400


def test_el_pdf_muestra_periodo_atajo_filtros_y_fechas(client, monkeypatch, iniciar_sesion, fake_connection_factory):
    iniciar_sesion()
    parametros = {"id_cliente": 10, "fecha_inicio": "2026-08-25", "fecha_fin": "2026-09-25",
                  "periodo_etiqueta": "mes"}
    conexion = fake_connection_factory(_handler_exportacion(parametros))
    monkeypatch.setattr(reportes_routes, "get_connection", lambda: conexion)

    respuesta = client.get("/api/reportes/5/pdf")
    texto = " ".join(PdfReader(io.BytesIO(respuesta.data)).pages[0].extract_text().split())

    assert respuesta.status_code == 200
    assert "Período: 2026-08-25 a 2026-09-25 (Último mes)" in texto
    assert "Cliente: Cliente Diez" in texto
    assert "Fecha de generación: 01-09-2026 08:30" in texto
    assert "Exportado:" in texto


def test_el_excel_conserva_la_hoja_de_datos_y_agrega_parametros(
    client, monkeypatch, iniciar_sesion, fake_connection_factory
):
    iniciar_sesion()
    parametros = {"id_cliente": 10, "fecha_inicio": "2026-08-25", "fecha_fin": "2026-09-25",
                  "periodo_etiqueta": "mes"}
    conexion = fake_connection_factory(_handler_exportacion(parametros))
    monkeypatch.setattr(reportes_routes, "get_connection", lambda: conexion)

    respuesta = client.get("/api/reportes/5/excel")
    libro = load_workbook(io.BytesIO(respuesta.data))

    assert libro.sheetnames == ["Reporte", "Parámetros"]
    assert len(list(libro["Reporte"].iter_rows(values_only=True))) == 2
    parametros_hoja = dict(libro["Parámetros"].iter_rows(values_only=True))
    assert parametros_hoja["Reporte"] == "Tareas por cliente"
    assert parametros_hoja["Período"] == "2026-08-25 a 2026-09-25 (Último mes)"
    assert parametros_hoja["Cliente"] == "Cliente Diez"
    assert parametros_hoja["Fecha de generación"] == "01-09-2026 08:30"
    assert parametros_hoja["Identificador del reporte"] == 5
    assert "2026-08-25_a_2026-09-25" in respuesta.headers["Content-Disposition"]


def test_el_excel_de_productividad_informa_el_area_filtrada(
    client, monkeypatch, iniciar_sesion, fake_connection_factory
):
    iniciar_sesion()
    parametros = {"fecha_inicio": "2026-09-01", "fecha_fin": "2026-09-25", "id_area": "2"}
    conexion = fake_connection_factory(_handler_exportacion(parametros, "PRODUCTIVIDAD_POR_USUARIO"))
    monkeypatch.setattr(reportes_routes, "get_connection", lambda: conexion)

    libro = load_workbook(io.BytesIO(client.get("/api/reportes/5/excel").data))
    parametros_hoja = dict(libro["Parámetros"].iter_rows(values_only=True))

    assert parametros_hoja["Área"] == "CONTABLE"
    assert parametros_hoja["Período"] == "2026-09-01 a 2026-09-25"
